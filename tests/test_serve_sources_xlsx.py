"""The source picker listing a cited XLSX (the picker half of the xlsx bullet
in docs/design/editor-source-picker.md §8; §2/§5 reads applied to a workbook).

Through the real HTTP surface against a real fixture project, in the posture
`test_serve_sources.py` uses: a real `EditorApp`, real workbook bytes written
by openpyxl, and a `snapshot_tree` comparison proving the reads write nothing.
The acceptance gate of `calc-sources.md` §11 question 2 is proven end to end
here: a key taken from the picker's listing composes its line through the
proposal read, and `refdes fetch --update` accepts the value that line names.
"""

from __future__ import annotations

import pytest
import yaml
from serve_support import Client, snapshot_tree
from test_serve_sources import make_root, sources_url, start

from refdes import cli as cli_mod

pytest.importorskip("openpyxl")

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

XLSX_ITEMS = """\
defaults: { type: decision }
items:
  - id: DEC-001
    citations:
      - path: analysis/budget.xlsx
    body: |
      ```calc
      P = source("analysis/budget.xlsx", "power") | W
      Q = 1 | 1
      ```
"""


def write_workbook(path, *, power=1.85):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["B2"] = power
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["C3"] = 2.3
    workbook.defined_names.add(DefinedName("power", attr_text="'Sheet1'!$B$2"))
    hidden.defined_names.add(DefinedName("local", attr_text="'Hidden'!$C$3"))
    workbook.save(path)
    workbook.close()


def make_xlsx_root(tmp_path, *, power=1.85):
    root = make_root(tmp_path, items=XLSX_ITEMS)
    write_workbook(root / "analysis" / "budget.xlsx", power=power)
    return root


def lock(root):
    return yaml.safe_load((root / ".refdes" / "citations.yaml").read_text("utf-8"))


# ------------------------------------------------------------------ 1: the file


def test_the_picker_lists_a_cited_xlsx_as_key_rows(tmp_path):
    root = make_xlsx_root(tmp_path)
    app = start(root)
    try:
        client = Client(app)
        status, payload = client.api_get(sources_url("DEC-001"))
        assert status == 200, payload
        assert payload["problems"] == []
        (budget,) = payload["files"]
        assert budget["path"] == "analysis/budget.xlsx"
        assert budget["reader"] == "xlsx"
        assert budget["browse"] == "rows"
        assert budget["state"] == "unpinned"
    finally:
        app.stop()


def test_listing_a_cited_xlsx_writes_nothing(tmp_path):
    root = make_xlsx_root(tmp_path)
    app = start(root)
    try:
        client = Client(app)
        before = snapshot_tree(root)
        status, payload = client.api_get(sources_url("DEC-001"))
        assert status == 200, payload
        status, payload = client.api_get(
            sources_url("DEC-001", "/entries") + "?path=analysis/budget.xlsx"
        )
        assert status == 200, payload
        status, payload = client.api_get(
            sources_url("DEC-001", "/propose")
            + "?path=analysis/budget.xlsx&key=power&unit=W"
        )
        assert status == 200, payload
        assert snapshot_tree(root) == before, "listing an xlsx wrote something"
    finally:
        app.stop()


# ------------------------------------------------------------------ 2: the rows


def test_the_entries_endpoint_lists_defined_names_with_qualified_keys(tmp_path):
    root = make_xlsx_root(tmp_path)
    app = start(root)
    try:
        client = Client(app)
        status, payload = client.api_get(
            sources_url("DEC-001", "/entries") + "?path=analysis/budget.xlsx"
        )
        assert status == 200, payload
        assert payload["reader"] == "xlsx"
        assert [e["key"] for e in payload["entries"]] == ["power", "Hidden!local"]
        power, local = payload["entries"]
        assert (power["raw"], power["value"], power["line"]) == ("1.85", "1.85", 2)
        assert power["context"] == [{"name": "cell", "text": "Sheet1!B2"}]
        assert power["selectable"] and power["problem"] == ""
        assert {
            c["name"]: c["text"] for c in local["context"]
        } == {"cell": "Hidden!C3", "sheet": "hidden"}
        assert payload["truncated"] is False
        assert payload["limits"]["max_rows"] == 5000
        assert payload["query"] == ""
        # The server-side filter is a display filter, as it is for CSV rows.
        status, filtered = client.api_get(
            sources_url("DEC-001", "/entries")
            + "?path=analysis/budget.xlsx&q=hidden"
        )
        assert status == 200
        assert [e["key"] for e in filtered["entries"]] == ["Hidden!local"]
        # No absolute server path leaves the process, in any string.
        for value in _strings(payload):
            assert str(root) not in value
    finally:
        app.stop()


def test_entries_show_the_pinned_value_and_flag_drift(tmp_path):
    root = make_xlsx_root(tmp_path)
    assert cli_mod.main([
        "-c", str(root / "refdes-project.yaml"), "fetch",
    ]) == 0
    workbook = load_workbook(root / "analysis" / "budget.xlsx")
    workbook["Sheet1"]["B2"] = 2.3
    workbook.save(root / "analysis" / "budget.xlsx")
    workbook.close()
    app = start(root)
    try:
        client = Client(app)
        status, payload = client.api_get(
            sources_url("DEC-001", "/entries") + "?path=analysis/budget.xlsx"
        )
        assert status == 200, payload
        power = next(e for e in payload["entries"] if e["key"] == "power")
        assert power["value"] == "2.3"  # what a pick would extract now
        assert power["pinned"] == "1.85"  # what the lockfile recorded
        assert power["changed"] is True
    finally:
        app.stop()


def test_the_entries_endpoint_reports_an_unreadable_workbook_as_a_panel_state(tmp_path):
    root = make_root(tmp_path, items=XLSX_ITEMS)
    (root / "analysis" / "budget.xlsx").write_bytes(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 512
    )
    app = start(root)
    try:
        client = Client(app)
        status, payload = client.api_get(
            sources_url("DEC-001", "/entries") + "?path=analysis/budget.xlsx"
        )
        assert status == 200, payload  # a panel state, not a request failure
        assert payload["entries"] == []
        assert any("encrypted" in problem for problem in payload["problems"])
        assert payload["problems"], payload
        for problem in payload["problems"]:
            assert str(root) not in problem
    finally:
        app.stop()


# --------------------------------------------------------------- 3: the line


def test_the_proposal_composes_a_line_for_a_picked_xlsx_key(tmp_path):
    root = make_xlsx_root(tmp_path)
    app = start(root)
    try:
        client = Client(app)
        status, payload = client.api_get(
            sources_url("DEC-001", "/propose")
            + "?path=analysis/budget.xlsx&key=power&unit=W"
        )
        assert status == 200, payload
        assert payload["line"] == 'power = source("analysis/budget.xlsx", "power") | W'
        assert payload["complete"] is True
        assert payload["entry"]["raw"] == "1.85"
        # The qualified sheet-scoped key round-trips through the composed line.
        status, qualified = client.api_get(
            sources_url("DEC-001", "/propose")
            + "?path=analysis/budget.xlsx&key=Hidden%21local&unit=W"
        )
        assert status == 200, qualified
        assert qualified["line"] == (
            'hidden_local = source("analysis/budget.xlsx", "Hidden!local") | W'
        )
        # With no unit there is no line, as for every keyed reader.
        status, bare = client.api_get(
            sources_url("DEC-001", "/propose") + "?path=analysis/budget.xlsx&key=power"
        )
        assert status == 200 and bare["line"] is None and bare["complete"] is False
    finally:
        app.stop()


# -------------------------------------------- the §11 Q2 acceptance gate, e2e


def test_a_picked_xlsx_key_fetches_and_fetch_update_is_the_acceptance_gate(
    tmp_path, capsys
):
    """The recorded acceptance gate (calc-sources.md §11 Q2): a key chosen from
    the picker's listing, composed by the proposal read, must work end to end
    through `refdes fetch --update` -- first pin, loud drift, accepted update
    -- with the workbook itself never written."""
    root = make_xlsx_root(tmp_path)
    app = start(root)
    try:
        client = Client(app)
        status, listing = client.api_get(
            sources_url("DEC-001", "/entries") + "?path=analysis/budget.xlsx"
        )
        assert status == 200 and listing["entries"][0]["key"] == "power"
        status, proposal = client.api_get(
            sources_url("DEC-001", "/propose")
            + "?path=analysis/budget.xlsx&key=power&unit=W"
        )
        assert status == 200 and proposal["complete"]
        line = proposal["line"]
    finally:
        app.stop()

    # The author inserts the composed line (the panel does this through a body
    # save; here the same file edit a save performs) and runs the gate command.
    items = root / "items" / "decisions.yaml"
    items.write_text(
        items.read_text("utf-8").replace(
            "Q = 1 | 1\n", f"Q = 1 | 1\n      {line}\n"
        ),
        encoding="utf-8",
    )
    workbook_bytes = (root / "analysis" / "budget.xlsx").read_bytes()
    assert cli_mod.main([
        "-c", str(root / "refdes-project.yaml"), "fetch", "--update",
    ]) == 0
    assert "extracted analysis/budget.xlsx power = 1.85" in capsys.readouterr().out
    assert (root / "analysis" / "budget.xlsx").read_bytes() == workbook_bytes
    record = lock(root)["citations"]["analysis/budget.xlsx"]
    assert record["values"] == {"power": {"reader": "xlsx", "value": "1.85"}}

    # Drift: a changed file warns and keeps the pin; fetch --update accepts it.
    workbook = load_workbook(root / "analysis" / "budget.xlsx")
    workbook["Sheet1"]["B2"] = 2.3
    workbook.save(root / "analysis" / "budget.xlsx")
    workbook.close()
    assert cli_mod.main([
        "-c", str(root / "refdes-project.yaml"), "fetch",
    ]) == 0
    capsys.readouterr()
    assert lock(root)["citations"]["analysis/budget.xlsx"]["values"] == record["values"]
    assert cli_mod.main([
        "-c", str(root / "refdes-project.yaml"), "fetch", "--update",
    ]) == 0
    assert "1.85 -> 2.3" in capsys.readouterr().out
    updated = lock(root)["citations"]["analysis/budget.xlsx"]
    assert updated["values"] == {"power": {"reader": "xlsx", "value": "2.3"}}
    assert updated["sha256"] != record["sha256"]


def _strings(payload):
    if isinstance(payload, str):
        yield payload
    elif isinstance(payload, dict):
        for value in payload.values():
            yield from _strings(value)
    elif isinstance(payload, list):
        for value in payload:
            yield from _strings(value)

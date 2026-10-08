"""XLSX defined names and the shared source fetch/lock contract."""

from __future__ import annotations

from decimal import Decimal

import pytest
import yaml
from conftest import write_project_config

from refdes import build, cli, keys, links, parse, sources
from refdes.schema import load_project
from refdes.sources import SourceExtractionError, SourceRequest


@pytest.fixture
def spreadsheet(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.workbook.defined_name import DefinedName

    path = tmp_path / "budget.xlsx"
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Sheet1"
    second = workbook.create_sheet("Sheet2")
    first["B2"] = 1.85
    second["B2"] = 2.3
    workbook.defined_names.add(DefinedName("power", attr_text="'Sheet1'!$B$2"))
    first.defined_names.add(DefinedName("local", attr_text="'Sheet1'!$B$2"))
    workbook.save(path)
    workbook.close()
    return path


def _add_name(path, name, destination, sheet=None):
    from openpyxl import load_workbook
    from openpyxl.workbook.defined_name import DefinedName

    workbook = load_workbook(path)
    owner = workbook.defined_names if sheet is None else workbook[sheet].defined_names
    owner.add(DefinedName(name, attr_text=destination))
    workbook.save(path)
    workbook.close()


def _extract(path, *names):
    return sources.reader_for("BUDGET.XLSX").extract(
        path, [SourceRequest("budget.xlsx", name) for name in names], label="budget.xlsx"
    )


def _fails(path, name, fragment):
    with pytest.raises(SourceExtractionError) as info:
        _extract(path, name)
    assert fragment in str(info.value), str(info.value)


def test_xlsx_workbook_and_sheet_scoped_names_return_cells(spreadsheet):
    got = _extract(spreadsheet, "power", "Sheet1!local", "local")
    assert {key: item.value for key, item in got.items()} == {
        "power": Decimal("1.85"),
        "Sheet1!local": Decimal("1.85"),
        "local": Decimal("1.85"),
    }
    assert {item.reader for item in got.values()} == {"xlsx"}


def test_xlsx_missing_name_and_missing_sheet_fail(spreadsheet):
    _fails(spreadsheet, "absent", "no defined name")
    _fails(spreadsheet, "Missing!local", "sheet 'Missing' does not exist")
    _add_name(spreadsheet, "orphan", "'Missing'!$A$1")
    _fails(spreadsheet, "orphan", "destination sheet 'Missing' does not exist")


def test_xlsx_sheet_scope_ambiguity_and_qualified_no_fallback(spreadsheet):
    _add_name(spreadsheet, "local", "'Sheet2'!$B$2", sheet="Sheet2")
    _fails(spreadsheet, "local", "ambiguous")
    with pytest.raises(SourceExtractionError) as info:
        _extract(spreadsheet, "local")
    assert "Sheet1" in str(info.value) and "Sheet2" in str(info.value)
    assert _extract(spreadsheet, "Sheet2!local")["Sheet2!local"].value == Decimal("2.3")
    _fails(spreadsheet, "Sheet2!power", "no defined name")


def test_xlsx_rejects_ranges_external_and_non_numeric_cells(spreadsheet):
    _add_name(spreadsheet, "wide", "'Sheet1'!$A$1:$B$2")
    _add_name(spreadsheet, "external", "'[other.xlsx]Sheet1'!$A$1")
    _fails(spreadsheet, "wide", "one cell")
    _fails(spreadsheet, "external", "external workbook")
    from openpyxl import load_workbook

    for value in ("3.2", True, "#DIV/0!", None):
        workbook = load_workbook(spreadsheet)
        workbook["Sheet1"]["B2"] = value
        workbook.save(spreadsheet)
        workbook.close()
        _fails(spreadsheet, "power", "blank" if value is None else "not a numeric cell")


def test_xlsx_formula_without_cached_value_has_actionable_error(spreadsheet):
    from openpyxl import load_workbook

    workbook = load_workbook(spreadsheet)
    workbook["Sheet1"]["B2"] = "=1+2"
    workbook.save(spreadsheet)
    workbook.close()
    _fails(spreadsheet, "power", "open the workbook in a calculating application")


def test_xlsx_missing_extra_names_install_command(tmp_path, monkeypatch, capsys):
    original_import = sources.importlib.import_module

    def missing_openpyxl(name, *args, **kwargs):
        if name == "openpyxl":
            raise ModuleNotFoundError("No module named 'openpyxl'", name="openpyxl")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(sources.importlib, "import_module", missing_openpyxl)
    _fails(tmp_path / "budget.xlsx", "power", 'pip install "refdes[xlsx]"')
    config = write_project_config(
        tmp_path,
        "site: {title: Test, out: _site}\n"
        "types:\n  decision:\n    prefix: DEC\n    fields:\n"
        "      citations: {type: citations, on_change: invalidate}\n",
    )
    (tmp_path / "budget.xlsx").write_bytes(b"not read without the extra")
    (tmp_path / "items").mkdir()
    (tmp_path / "items" / "p.md").write_text(
        "---\nid: DEC-001\ntype: decision\ncitations:\n  - path: budget.xlsx\n---\n\n"
        '```calc\nP = source("budget.xlsx", "power") | W\n```\n',
        encoding="utf-8",
    )
    assert cli.main(["-c", str(config), "fetch"]) == 1
    assert 'pip install "refdes[xlsx]"' in capsys.readouterr().err


def _project(config):
    project = load_project(config_path=str(config))
    parse.load_items(project)
    if keys.mint_missing(project, write=True):
        project.items = {}
        project.items_by_id = {}
        project.pending = []
        parse.load_items(project)
    links.expand_missing_calc_refs(project, write=True)
    build.build(project)
    assert not project.errors, [d.message for d in project.errors]
    return project


def _record(tmp_path):
    lock = yaml.safe_load((tmp_path / ".refdes" / "citations.yaml").read_text("utf-8"))
    return lock["citations"]["budget.xlsx"]


@pytest.mark.parametrize("key", ["power", "Sheet1!local", "local"])
def test_xlsx_fetch_pins_value_and_only_accepted_cell_edit_changes_hash(
    spreadsheet, tmp_path, capsys, key
):
    config = write_project_config(
        tmp_path,
        "site: {title: Test, out: _site}\n"
        "types:\n  decision:\n    prefix: DEC\n    fields:\n"
        "      citations: {type: citations, on_change: invalidate}\n",
    )
    (tmp_path / "items").mkdir()
    (tmp_path / "items" / "p.md").write_text(
        "---\nid: DEC-001\ntype: decision\ncitations:\n  - path: budget.xlsx\n---\n\n"
        f'```calc\nP = source("budget.xlsx", "{key}") | W\n```\n',
        encoding="utf-8",
    )
    initial_bytes = spreadsheet.read_bytes()
    assert cli.main(["-c", str(config), "fetch"]) == 0
    assert spreadsheet.read_bytes() == initial_bytes
    assert f"extracted budget.xlsx {key} = 1.85" in capsys.readouterr().out
    first = _record(tmp_path)
    assert first["values"] == {key: {"reader": "xlsx", "value": "1.85"}}
    first_hash = _project(config).item_by_id("DEC-001").content_hash

    from openpyxl import load_workbook

    workbook = load_workbook(spreadsheet)
    workbook["Sheet1"]["B2"] = 2.3
    workbook.save(spreadsheet)
    workbook.close()
    edited_bytes = spreadsheet.read_bytes()
    assert _project(config).item_by_id("DEC-001").content_hash == first_hash
    assert cli.main(["-c", str(config), "fetch"]) == 0
    assert _record(tmp_path) == first
    assert cli.main(["-c", str(config), "fetch", "--update"]) == 0
    assert spreadsheet.read_bytes() == edited_bytes
    second = _record(tmp_path)
    assert second["sha256"] != first["sha256"]
    assert second["values"] == {key: {"reader": "xlsx", "value": "2.3"}}
    assert _project(config).item_by_id("DEC-001").content_hash != first_hash
    assert _project(config).item_by_id("DEC-001").calc_values["P"] == "2.3 W"
